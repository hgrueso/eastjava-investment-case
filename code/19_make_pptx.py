# 19_make_pptx.py: PowerPoint version of the deck for comments.
# Each PDF page becomes a full-slide image (layout identical to the PDF), and
# the speaker notes written in the .qmd are copied into each slide's notes field.
# Checks that the PDF is the current version before building, and saves with a
# timestamp so PowerPoint cannot show a cached older copy.
import os, re, html, glob, shutil, tempfile, datetime
import pymupdf as fitz
from pptx import Presentation
from pptx.util import Emu

PDF, HTML = "output/06_slides_en.pdf", "code/06_slides_en.html"
pdf = fitz.open(PDF)
mtime = datetime.datetime.fromtimestamp(os.path.getmtime(PDF)).strftime("%Y-%m-%d %H:%M")
labels = sum(page.get_text().count("Data source:") for page in pdf)
print(f"PDF: {PDF} | {len(pdf)} pages | last modified {mtime} | 'Data source' labels found: {labels}")
if labels == 0:
    print("STOP: this PDF is an older version (no 'Data source' labels). Re-render the deck before building the PowerPoint.")
    raise SystemExit(1)

H = open(HTML, encoding="utf-8").read()
body = H[H.find('<div class="slides">'):]
notes = []
for part in re.split(r'(?=<section\b)', body)[1:]:
    tag = part[:part.find('>') + 1]
    if 'id="title-slide"' in tag or re.search(r'class="[^"]*\bslide\b', tag):
        m = re.search(r'<aside class="notes">(.*?)</aside>', part, re.S)
        notes.append(re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', ' ', m.group(1)))).strip() if m else "")
if len(notes) != len(pdf):
    print(f"WARNING: {len(notes)} slides in the HTML but {len(pdf)} PDF pages; notes skipped to avoid misalignment.")
    notes = [""] * len(pdf)

tmp = tempfile.mkdtemp()
prs = Presentation()
w, h = pdf[0].rect.width, pdf[0].rect.height
prs.slide_width = Emu(12192000); prs.slide_height = Emu(int(12192000 * h / w))
for i, page in enumerate(pdf):
    f = os.path.join(tmp, f"slide_{i:02d}.png"); page.get_pixmap(dpi=200).save(f)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(f, 0, 0, prs.slide_width, prs.slide_height)
    if notes[i]: s.notes_slide.notes_text_frame.text = notes[i]
shutil.rmtree(tmp)

for old in glob.glob("output/06_slides_en*.pptx"): os.remove(old)
stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
out = f"output/06_slides_en_{stamp}.pptx"
prs.save(out)
print(f"Saved {len(pdf)} slides, {sum(1 for n in notes if n)} with speaker notes -> {out}")
